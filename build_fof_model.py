"""
Builds the Fund of Funds Fee Model Excel workbook.
Run: python build_fof_model.py  →  fof_fee_model.xlsx
"""

import openpyxl
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side,
)
from openpyxl.utils import get_column_letter

# ── Palette ────────────────────────────────────────────────────────────────────
NAVY        = "1F3864"
HEADER_BLUE = "2F5496"
MID_BLUE    = "4472C4"
PALE_BLUE   = "DAE3F3"
INPUT_AMBER = "FFF2CC"
CALC_GRAY   = "F2F2F2"
WHITE       = "FFFFFF"
BORDER_GRAY = "BFBFBF"
GREEN_DARK  = "375623"
GREEN_LIGHT = "E2EFDA"
RED_DARK    = "9C0006"
RED_LIGHT   = "FFC7CE"

# ── Cell reference anchors (all formulas reference these) ──────────────────────
# Input cells (column C)
R_FUND_SIZE   = 5   # C5  – Fund Size ($M)
R_FUND_LIFE   = 6   # C6  – Fund Life (years)
R_MGMT_RATE   = 9   # C9  – Mgmt Fee % p.a.
R_MGMT_TOTAL  = 10  # C10 – Total Mgmt Fees ($M) [derived]
R_INVESTED    = 11  # C11 – Invested Capital ($M) [derived]
R_PLACE_RATE  = 14  # C14 – Placement Fee %
R_HURDLE      = 18  # C18 – Preferred Return % p.a.
R_PREDPI_GP   = 19  # C19 – Pre-DPI GP Carry %
R_DPI_THRESH  = 20  # C20 – DPI Threshold (x)
R_CATCHUP_TGT = 21  # C21 – Catch-up Target GP %
R_POST_GP     = 22  # C22 – Post Catch-up GP %
R_PREF_AMT    = 25  # C25 – Preferred Return $M [derived]
R_DPI_AMT     = 26  # C26 – LP DPI Threshold $M [derived]

INPUT_COL  = "C"   # column for input values
LABEL_COL  = "B"   # column for row labels

# Scenario table
SCENARIO_HDR_ROW = 30
SCENARIO_START   = 31   # first data row
SCENARIO_END     = 37   # last data row  (1x – 7x)
GROSS_MULTIPLES  = list(range(1, 8))

# Visible columns A-I, hidden working columns J-AB
VIS_COLS = {
    "A": ("Gross\nMOIC",         12),
    "B": ("Gross Proceeds\n($M)", 17),
    "C": ("Net MOIC\n(LP)",       13),
    "D": ("LP Proceeds\n($M)",    17),
    "E": ("Mgmt Fee\n($M)",       14),
    "F": ("Placement\nFee ($M)",  14),
    "G": ("GP Carry\n($M)",       14),
    "H": ("Total Fees\n($M)",     14),
    "I": ("Regime",               26),
}

# Working column letters (J=10 … AB=28)
def wc(offset):
    """Return column letter for a working column (offset 0 = J)."""
    return get_column_letter(10 + offset)

WC = {name: wc(i) for i, name in enumerate([
    "gross_profit",   # J
    "placement",      # K
    "net_to_wfall",   # L
    "cap_returned",   # M
    "net_aft_cap",    # N
    "pref_paid",      # O
    "net_aft_pref",   # P
    "lp_to_dpi",      # Q
    "predpi_pool",    # R
    "predpi_amt",     # S
    "predpi_lp",      # T
    "predpi_gp",      # U
    "rem_predpi",     # V
    "gp_target",      # W
    "gp_needed",      # X
    "catchup_amt",    # Y
    "rem_catchup",    # Z
    "post_lp",        # AA
    "post_gp",        # AB
])}

# ── Style helpers ───────────────────────────────────────────────────────────────
def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def font(bold=False, color="000000", size=11, italic=False):
    return Font(bold=bold, color=color, size=size, italic=italic,
                name="Calibri")

def align(h="left", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def thin_border(top=False, bottom=False, left=False, right=False):
    s = Side(style="thin", color=BORDER_GRAY)
    n = None
    return Border(
        top=s if top else n,
        bottom=s if bottom else n,
        left=s if left else n,
        right=s if right else n,
    )

def thick_border_bottom():
    return Border(bottom=Side(style="medium", color=NAVY))

def apply_row_border(ws, row, col_start, col_end, **kwargs):
    b = thin_border(**kwargs)
    for c in range(col_start, col_end + 1):
        ws.cell(row=row, column=c).border = b


def style_input(cell):
    cell.fill      = fill(INPUT_AMBER)
    cell.font      = font(bold=True, size=11)
    cell.alignment = align("right")
    cell.border    = thin_border(top=True, bottom=True, left=True, right=True)

def style_calc(cell):
    cell.fill      = fill(CALC_GRAY)
    cell.font      = font(italic=True, color="404040")
    cell.alignment = align("right")
    cell.border    = thin_border(top=True, bottom=True, left=True, right=True)

def style_section_header(ws, row, col_start, col_end, label, bg=HEADER_BLUE):
    ws.merge_cells(
        start_row=row, start_column=col_start,
        end_row=row,   end_column=col_end
    )
    c = ws.cell(row=row, column=col_start, value=label)
    c.fill      = fill(bg)
    c.font      = font(bold=True, color=WHITE, size=10)
    c.alignment = align("left", "center")
    ws.row_dimensions[row].height = 16


# ── Build workbook ─────────────────────────────────────────────────────────────
def build():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws.sheet_view.showGridLines = False

    # ── Column widths ──────────────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 3    # left margin
    ws.column_dimensions[LABEL_COL].width = 32
    ws.column_dimensions[INPUT_COL].width = 16
    ws.column_dimensions["D"].width = 3    # gap

    for col_letter, (_, w) in VIS_COLS.items():
        ws.column_dimensions[col_letter].width = w

    for name, col_letter in WC.items():
        ws.column_dimensions[col_letter].width = 14
        ws.column_dimensions[col_letter].hidden = True

    # ── Title ──────────────────────────────────────────────────────────────────
    ws.row_dimensions[1].height = 10
    ws.row_dimensions[2].height = 32
    ws.merge_cells("B2:I2")
    t = ws["B2"]
    t.value     = "FUND OF FUNDS  —  FEE STRUCTURE MODEL"
    t.fill      = fill(NAVY)
    t.font      = font(bold=True, color=WHITE, size=16)
    t.alignment = align("left", "center")

    ws.row_dimensions[3].height = 10

    # ── INPUTS block (rows 4 – 27) ─────────────────────────────────────────────
    def label_row(row, text, height=18):
        c = ws.cell(row=row, column=2, value=text)
        c.font      = font(size=10)
        c.alignment = align("left", "center")
        ws.row_dimensions[row].height = height

    def input_row(row, label, value, fmt, note=""):
        label_row(row, label)
        c = ws.cell(row=row, column=3, value=value)
        c.number_format = fmt
        style_input(c)
        if note:
            n = ws.cell(row=row, column=4, value=note)
            n.font      = font(italic=True, color="808080", size=9)
            n.alignment = align("left", "center")

    def calc_row(row, label, formula, fmt):
        label_row(row, label)
        c = ws.cell(row=row, column=3, value=formula)
        c.number_format = fmt
        style_calc(c)

    # ── Section: Fund Parameters ───────────────────────────────────────────────
    style_section_header(ws, 4,  2, 3, "  FUND PARAMETERS")
    input_row(R_FUND_SIZE,  "Fund Size ($M)",    30,  '#,##0.0')
    input_row(R_FUND_LIFE,  "Fund Life (years)",  10, '0')

    # ── Section: Management Fee ────────────────────────────────────────────────
    ws.row_dimensions[7].height = 8
    style_section_header(ws, 8, 2, 3, "  MANAGEMENT FEE")
    input_row(R_MGMT_RATE,  "Rate (% p.a.)",     0.015, '0.0%')
    calc_row(R_MGMT_TOTAL,
             "Total Mgmt Fees ($M)",
             f"={INPUT_COL}{R_FUND_SIZE}*{INPUT_COL}{R_MGMT_RATE}*{INPUT_COL}{R_FUND_LIFE}",
             '#,##0.0')
    calc_row(R_INVESTED,
             "Invested Capital ($M)",
             f"={INPUT_COL}{R_FUND_SIZE}-{INPUT_COL}{R_MGMT_TOTAL}",
             '#,##0.0')

    # ── Section: Placement Fee ─────────────────────────────────────────────────
    ws.row_dimensions[12].height = 8
    style_section_header(ws, 13, 2, 3, "  PLACEMENT FEE  (paid to underlying managers)")
    input_row(R_PLACE_RATE,
              "Rate (% of gross portfolio profit)", 0.10, '0.0%',
              note="← no hurdle; applied from first dollar of profit")

    # ── Section: Carry Waterfall ───────────────────────────────────────────────
    ws.row_dimensions[15].height = 8
    style_section_header(ws, 16, 2, 3, "  CARRY WATERFALL")
    ws.row_dimensions[17].height = 8
    style_section_header(ws, 17, 2, 3,
                         "  (1) Return of Capital  →  (2) Preferred Return  →  (3) Pre-DPI Split  "
                         "→  (4) Catch-up  →  (5) Post Catch-up",
                         bg=MID_BLUE)
    input_row(R_HURDLE,
              "Preferred Return Hurdle (% p.a., compound)", 0.00, '0.0%',
              note="← set to 0% to disable")
    input_row(R_PREDPI_GP,  "Pre-DPI GP Carry (%)",           0.10, '0.0%')
    input_row(R_DPI_THRESH, "DPI Threshold (x)",              3,    '0.0"x"')
    input_row(R_CATCHUP_TGT,"Catch-up Target GP (%)",         0.15, '0.0%')
    input_row(R_POST_GP,    "Post Catch-up GP (%)",           0.15, '0.0%')

    # ── Section: Derived ──────────────────────────────────────────────────────
    ws.row_dimensions[23].height = 8
    style_section_header(ws, 24, 2, 3, "  DERIVED")
    calc_row(R_PREF_AMT,
             "Preferred Return ($M)",
             f"={INPUT_COL}{R_FUND_SIZE}*((1+{INPUT_COL}{R_HURDLE})^{INPUT_COL}{R_FUND_LIFE}-1)",
             '#,##0.0')
    calc_row(R_DPI_AMT,
             "LP DPI Threshold ($M)",
             f"={INPUT_COL}{R_FUND_SIZE}*{INPUT_COL}{R_DPI_THRESH}",
             '#,##0.0')

    # ── Scenario Table header ─────────────────────────────────────────────────
    ws.row_dimensions[28].height = 10
    ws.row_dimensions[29].height = 10
    style_section_header(ws, 29, 1, 9,
                         "  SCENARIO ANALYSIS  —  Gross MOIC on committed capital  "
                         "(change values in column A to flex scenarios)",
                         bg=NAVY)

    # Column headers (row 30)
    ws.row_dimensions[SCENARIO_HDR_ROW].height = 34
    for col_letter, (header, _) in VIS_COLS.items():
        col_num = openpyxl.utils.column_index_from_string(col_letter)
        c = ws.cell(row=SCENARIO_HDR_ROW, column=col_num, value=header)
        c.fill      = fill(PALE_BLUE)
        c.font      = font(bold=True, color=NAVY, size=10)
        c.alignment = align("center", "center", wrap=True)
        c.border    = thin_border(top=True, bottom=True, left=True, right=True)

    # ── Scenario data rows ────────────────────────────────────────────────────
    # Named absolute refs (for readability in formulas)
    FS  = f"$C${R_FUND_SIZE}"    # fund size
    INV = f"$C${R_INVESTED}"     # invested capital
    PR  = f"$C${R_PLACE_RATE}"   # placement rate
    HR  = f"$C${R_HURDLE}"       # hurdle rate
    FL  = f"$C${R_FUND_LIFE}"    # fund life
    PGP = f"$C${R_PREDPI_GP}"    # pre-DPI GP %
    CUP = f"$C${R_CATCHUP_TGT}"  # catchup target GP %
    POP = f"$C${R_POST_GP}"      # post-catchup GP %
    PRA = f"$C${R_PREF_AMT}"     # pref return $M
    DPA = f"$C${R_DPI_AMT}"      # DPI threshold $M
    MGT = f"$C${R_MGMT_TOTAL}"   # total mgmt fee

    def wcf(name, row):
        """Absolute-column, relative-row reference for a working column."""
        return f"${WC[name]}${row}"

    alt_fill   = fill("F5F8FF")
    plain_fill = fill(WHITE)

    for i, gm in enumerate(GROSS_MULTIPLES):
        row = SCENARIO_START + i
        bg  = alt_fill if i % 2 else plain_fill
        ws.row_dimensions[row].height = 18

        # ── Visible columns ──────────────────────────────────────────────────
        # A: Gross MOIC (editable input)
        ca = ws.cell(row=row, column=1, value=gm)
        ca.fill      = fill(INPUT_AMBER)
        ca.font      = font(bold=True, size=11)
        ca.alignment = align("center", "center")
        ca.number_format = '0"x"'
        ca.border    = thin_border(top=True, bottom=True, left=True, right=True)

        A = f"$A${row}"  # gross MOIC ref

        # B: Gross Proceeds
        cb = ws.cell(row=row, column=2,
                     value=f"={A}*{FS}")
        cb.number_format = '#,##0.0'
        cb.fill = bg; cb.font = font(size=11)
        cb.alignment = align("right", "center")
        cb.border = thin_border(top=True, bottom=True, left=True, right=True)

        B = f"$B${row}"

        # Working columns ────────────────────────────────────────────────────
        # J – gross portfolio profit
        ws.cell(row=row, column=openpyxl.utils.column_index_from_string(WC["gross_profit"]),
                value=f"=MAX(0,{B}-{INV})")
        J = wcf("gross_profit", row)

        # K – placement fee
        ws.cell(row=row, column=openpyxl.utils.column_index_from_string(WC["placement"]),
                value=f"={PR}*{J}")
        K = wcf("placement", row)

        # L – net to waterfall
        ws.cell(row=row, column=openpyxl.utils.column_index_from_string(WC["net_to_wfall"]),
                value=f"={B}-{K}")
        L = wcf("net_to_wfall", row)

        # M – capital returned to LP
        ws.cell(row=row, column=openpyxl.utils.column_index_from_string(WC["cap_returned"]),
                value=f"=MIN({L},{FS})")
        M = wcf("cap_returned", row)

        # N – net after capital return
        ws.cell(row=row, column=openpyxl.utils.column_index_from_string(WC["net_aft_cap"]),
                value=f"=MAX(0,{L}-{FS})")
        N = wcf("net_aft_cap", row)

        # O – preferred return paid
        ws.cell(row=row, column=openpyxl.utils.column_index_from_string(WC["pref_paid"]),
                value=f"=MIN({N},{PRA})")
        O = wcf("pref_paid", row)

        # P – net after preferred return
        ws.cell(row=row, column=openpyxl.utils.column_index_from_string(WC["net_aft_pref"]),
                value=f"=MAX(0,{N}-{O})")
        P = wcf("net_aft_pref", row)

        # Q – LP still needed to reach DPI threshold (after capital + pref)
        ws.cell(row=row, column=openpyxl.utils.column_index_from_string(WC["lp_to_dpi"]),
                value=f"=MAX(0,{DPA}-{M}-{O})")
        Q = wcf("lp_to_dpi", row)

        # R – pre-DPI profit pool needed for LP to reach DPI threshold
        #     LP gets (1-PGP)% of this pool, so pool = Q / (1-PGP)
        ws.cell(row=row, column=openpyxl.utils.column_index_from_string(WC["predpi_pool"]),
                value=f"=IF({PGP}<1,{Q}/(1-{PGP}),0)")
        R = wcf("predpi_pool", row)

        # S – actual pre-DPI stage amount (capped at available profit)
        ws.cell(row=row, column=openpyxl.utils.column_index_from_string(WC["predpi_amt"]),
                value=f"=MIN({P},{R})")
        S = wcf("predpi_amt", row)

        # T – LP profit from pre-DPI stage
        ws.cell(row=row, column=openpyxl.utils.column_index_from_string(WC["predpi_lp"]),
                value=f"=(1-{PGP})*{S}")
        T = wcf("predpi_lp", row)

        # U – GP carry from pre-DPI stage
        ws.cell(row=row, column=openpyxl.utils.column_index_from_string(WC["predpi_gp"]),
                value=f"={PGP}*{S}")
        U = wcf("predpi_gp", row)

        # V – remaining profit after pre-DPI stage
        ws.cell(row=row, column=openpyxl.utils.column_index_from_string(WC["rem_predpi"]),
                value=f"=MAX(0,{P}-{S})")
        V = wcf("rem_predpi", row)

        # W – GP total carry target (catchup% of ALL profit above capital incl pref)
        ws.cell(row=row, column=openpyxl.utils.column_index_from_string(WC["gp_target"]),
                value=f"={CUP}*{N}")
        W = wcf("gp_target", row)

        # X – additional GP carry needed to hit target
        ws.cell(row=row, column=openpyxl.utils.column_index_from_string(WC["gp_needed"]),
                value=f"=MAX(0,{W}-{U})")
        X = wcf("gp_needed", row)

        # Y – catch-up amount actually paid
        ws.cell(row=row, column=openpyxl.utils.column_index_from_string(WC["catchup_amt"]),
                value=f"=MIN({V},{X})")
        Y = wcf("catchup_amt", row)

        # Z – remaining after catch-up
        ws.cell(row=row, column=openpyxl.utils.column_index_from_string(WC["rem_catchup"]),
                value=f"=MAX(0,{V}-{Y})")
        Z = wcf("rem_catchup", row)

        # AA – post catch-up LP share
        ws.cell(row=row, column=openpyxl.utils.column_index_from_string(WC["post_lp"]),
                value=f"=(1-{POP})*{Z}")
        AA = wcf("post_lp", row)

        # AB – post catch-up GP share
        ws.cell(row=row, column=openpyxl.utils.column_index_from_string(WC["post_gp"]),
                value=f"={POP}*{Z}")
        AB = wcf("post_gp", row)

        # ── Back to visible columns ──────────────────────────────────────────
        # D – LP Proceeds: capital + pref + pre-DPI LP profit + post-catchup LP
        cd = ws.cell(row=row, column=4,
                     value=f"={M}+{O}+{T}+{AA}")
        cd.number_format = '#,##0.0'
        cd.fill = bg; cd.font = font(bold=True, size=11)
        cd.alignment = align("right", "center")
        cd.border = thin_border(top=True, bottom=True, left=True, right=True)

        D = f"$D${row}"

        # C – Net MOIC to LP
        cc = ws.cell(row=row, column=3,
                     value=f"={D}/{FS}")
        cc.number_format = '0.00"x"'
        cc.fill = bg; cc.font = font(bold=True, size=11)
        cc.alignment = align("center", "center")
        cc.border = thin_border(top=True, bottom=True, left=True, right=True)

        # E – Mgmt Fees
        ce = ws.cell(row=row, column=5, value=f"={MGT}")
        ce.number_format = '#,##0.0'
        ce.fill = bg; ce.font = font(size=11)
        ce.alignment = align("right", "center")
        ce.border = thin_border(top=True, bottom=True, left=True, right=True)

        # F – Placement Fee
        cf = ws.cell(row=row, column=6, value=f"={K}")
        cf.number_format = '#,##0.0'
        cf.fill = bg; cf.font = font(size=11)
        cf.alignment = align("right", "center")
        cf.border = thin_border(top=True, bottom=True, left=True, right=True)

        # G – GP Carry
        cg = ws.cell(row=row, column=7,
                     value=f"={U}+{Y}+{AB}")
        cg.number_format = '#,##0.0'
        cg.fill = bg; cg.font = font(size=11)
        cg.alignment = align("right", "center")
        cg.border = thin_border(top=True, bottom=True, left=True, right=True)

        G = f"$G${row}"

        # H – Total Fees
        ch = ws.cell(row=row, column=8,
                     value=f"=${ce.column_letter}{row}+${cf.column_letter}{row}+{G}")
        ch.number_format = '#,##0.0'
        ch.fill = bg; ch.font = font(bold=True, size=11, color=RED_DARK)
        ch.alignment = align("right", "center")
        ch.border = thin_border(top=True, bottom=True, left=True, right=True)

        # I – Regime label
        pre_lp_pct  = f"TEXT((1-{PGP}),\"0%\")"
        pre_gp_pct  = f"TEXT({PGP},\"0%\")"
        post_lp_pct = f"TEXT((1-{POP}),\"0%\")"
        post_gp_pct = f"TEXT({POP},\"0%\")"

        regime_formula = (
            f'=IF({L}<{FS},"Loss / below capital",'
            f'IF({P}<={R},"Pre-DPI ("&{pre_lp_pct}&"/"&{pre_gp_pct}&")",'
            f'IF({V}<{X},"Catch-up (in progress)",'
            f'"Post catch-up ("&{post_lp_pct}&"/"&{post_gp_pct}&")")))'
        )
        ci = ws.cell(row=row, column=9, value=regime_formula)
        ci.fill = bg; ci.font = font(italic=True, size=10, color="404040")
        ci.alignment = align("left", "center")
        ci.border = thin_border(top=True, bottom=True, left=True, right=True)

    # ── Bottom border on last data row ─────────────────────────────────────────
    for col in range(1, 10):
        c = ws.cell(row=SCENARIO_END, column=col)
        c.border = thin_border(top=True, bottom=True, left=True, right=True)

    # ── Fee summary totals row ────────────────────────────────────────────────
    TOTAL_ROW = SCENARIO_END + 1
    ws.row_dimensions[TOTAL_ROW].height = 4   # thin divider spacer

    # ── Notes section ─────────────────────────────────────────────────────────
    NOTE_ROW = SCENARIO_END + 2
    ws.row_dimensions[NOTE_ROW].height = 10
    ws.merge_cells(f"A{NOTE_ROW}:I{NOTE_ROW}")
    style_section_header(ws, NOTE_ROW, 1, 9, "  NOTES", bg=MID_BLUE)

    notes = [
        "Gross MOIC: measured on committed capital.  LP returns also on committed capital.",
        f"Mgmt Fee: 1.5% p.a. × fund life, charged on committed capital.  Invested = Committed − Mgmt Fees.",
        "Placement Fee: taken off gross portfolio profit (Gross Proceeds − Invested Capital) before the carry waterfall.",
        "Preferred Return: compound annual hurdle on committed capital.  LP receives full pref before any carry splits.",
        "Catch-up base: GP target carry % applied to ALL profit above return of capital (including the pref return amount).",
        "DPI Threshold: once LP cumulative distributions (capital + pref + pre-DPI profit) reach Threshold × Committed,",
        "  the catch-up kicks in.  During catch-up, 100% flows to GP until they reach the Catch-up Target GP %.",
        "Yellow cells  = editable inputs.  Grey cells = derived.  Working columns (J onward) are hidden.",
    ]
    for j, note in enumerate(notes):
        r = NOTE_ROW + 1 + j
        ws.row_dimensions[r].height = 15
        ws.merge_cells(f"A{r}:I{r}")
        c = ws.cell(row=r, column=1, value=note)
        c.font      = font(size=9, color="404040",
                           italic=note.startswith("  "))
        c.alignment = align("left", "center")
        c.fill      = fill("FAFAFA")

    # ── Freeze panes ──────────────────────────────────────────────────────────
    ws.freeze_panes = f"B{SCENARIO_START}"

    # ── Print / page setup ───────────────────────────────────────────────────
    ws.page_setup.orientation  = "landscape"
    ws.page_setup.fitToPage    = True
    ws.page_setup.fitToWidth   = 1
    ws.page_setup.fitToHeight  = 0
    ws.print_area = f"A1:I{NOTE_ROW + len(notes)}"

    # ── Save ──────────────────────────────────────────────────────────────────
    out = "/home/user/Main/fof_fee_model.xlsx"
    wb.save(out)
    print(f"Saved → {out}")
    return out


if __name__ == "__main__":
    build()
